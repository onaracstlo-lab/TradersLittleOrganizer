"""Build 522 packaging/build-system remediation regressions."""
from tests import _release_artifacts as RA
from pathlib import Path
import struct
from docx import Document

import pytest

pytestmark = pytest.mark.behavior
ROOT = Path(__file__).resolve().parents[2]


def _text(name):
    return (ROOT / name).read_text(encoding="utf-8", errors="ignore")


def test_build522_native_dist_roots_derive_public_version():
    win = _text("createWindowsDist.ps1")
    linux = _text("createLinuxDist.sh")
    mac = _text("createMacOSDist.sh")
    assert "tloDist-V1.6" not in win + linux + mac
    assert "PUBLIC_VERSION" in win
    assert "PUBLIC_VERSION" in linux
    assert "PUBLIC_VERSION" in mac
    assert 'tloDist-V${PublicVersion}Build$BundleNumber' in win
    assert 'tloDist-V${PUBLIC_VERSION}Build${BUNDLE_NUMBER}' in linux
    assert 'tloDist-V${PUBLIC_VERSION}Build${BUNDLE_NUMBER}' in mac


def test_build522_main_icon_names_are_canonical_and_old_duplicates_removed():
    for suffix in ("png", "ico", "icns"):
        assert (ROOT / "icons" / f"tlo-main-icon.{suffix}").is_file()
        assert not (ROOT / "icons" / f"tlo-inventory-icon.{suffix}").exists()
    data = (ROOT / "icons" / "tlo-main-icon.ico").read_bytes()
    reserved, icon_type, count = struct.unpack_from("<HHH", data, 0)
    assert (reserved, icon_type) == (0, 1)
    assert count >= 5
    win = _text("createWindowsDist.ps1")
    mac = _text("createMacOSDist.sh")
    assert "tlo-main-icon.ico" in win and "tlo-inventory-icon.ico" not in win
    assert "tlo-main-icon.icns" in mac and "tlo-inventory-icon.icns" not in mac


def test_build522_macos_icons_are_required_and_pyinstaller7_is_rejected():
    mac = _text("createMacOSDist.sh")
    assert "find_required_icon" in mac
    assert "find_optional_icon" not in mac
    assert "Required macOS icon file not found" in mac
    assert "verify_build_environment.py" in mac
    verifier = _text("verify_build_environment.py")
    assert '"PyInstaller": "6.22.3"' in verifier


def test_build522_checked_in_ruff_contract_covers_review_hygiene_findings():
    config = _text("ruff.toml")
    for code in ("F401", "F541", "F821", "F822", "F823", "F841"):
        assert code in config
    phase = _text("tlo_phase23_v2.py")
    assert "tag_group_with_record, merge_tag_stats, emit_tag_fallback_summary" not in phase


def test_build522_docs_state_release_layout_and_macos_permission_consequence():
    manual = _text(RA.MANUAL_FILENAME)
    faq = _text("TLO-FAQ.txt")
    doc = Document(ROOT / RA.REQUIREMENTS_FILENAME)
    req = "\n".join([p.text for p in doc.paragraphs] + [cell.text for table in doc.tables for row in table.rows for cell in row.cells])
    assert "TLO_V1.7Build<BUILD>" in manual
    assert "Windows hybrid package uses seven PyInstaller onefile executables" in manual
    assert "tlo-main.app" in manual and "tlo-search.app" in manual and "search-artist-db.app" in manual
    assert "Full Disk Access" in manual and "Full Disk Access" in faq
    assert "Official release package layout" in req
    assert "tlo-deleteDupes.cmd" in req and "shared _internal" in req
    assert "PyInstaller < 7" in req
