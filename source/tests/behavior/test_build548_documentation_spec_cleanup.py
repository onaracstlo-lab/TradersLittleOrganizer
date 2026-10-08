"""Build 548: current documentation/specification and release-test metadata cleanup."""

from __future__ import annotations
from tests._release_artifacts import release_history

import ast
import re
from pathlib import Path

import pytest
from docx import Document

from tests import _release_artifacts as RA

pytestmark = pytest.mark.behavior
ROOT = Path(__file__).resolve().parents[2]


def _requirements_document() -> Document:
    return Document(ROOT / RA.REQUIREMENTS_FILENAME)


def _requirements_text() -> str:
    doc = _requirements_document()
    parts = [paragraph.text for paragraph in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def test_build548_version_and_current_artifact_contract():
    for filename in (
        RA.REQUIREMENTS_FILENAME,
        RA.MANUAL_FILENAME,
        RA.SOURCE_README_FILENAME,
        RA.CHANGES_FILENAME,
        RA.BUILD_VERIFICATION_FILENAME,
    ):
        assert (ROOT / filename).is_file()


def test_build548_requirements_corrects_cli_doc_and_revision_index_contracts():
    doc = _requirements_document()
    text = _requirements_text()
    heading = next(p for p in doc.paragraphs if p.text.strip().startswith("13.1 Standalone Tagger CLI Requirements"))
    assert heading.style.name == "Heading 2"
    assert "--tag-copy-delete-original is a tlo-main GUI-launcher flag" in text
    assert "Build 548: docs/spec cleanup" in text
    assert "Build 536: update-package safety -> §2.12." in text
    assert "Legacy binary .doc filenames may be recognized for diagnostics" in text
    assert ".doc content is unsupported and is not decoded as readable setlist text" in text


def test_build548_requirements_makes_ia05_traceable_without_claiming_schema2_runtime():
    text = _requirements_text()
    assert "B1. Open audit findings" in text
    assert "IA-05" in text
    assert "First-install/bootstrap authenticity and update-key rotation" in text
    assert "issued_utc" in text
    assert "expires_utc" in text
    assert "minimum_acceptable_build" in text
    assert "it does not change the updater or schema-1 trust semantics" in text


def test_build548_manual_warns_about_corruption_defaults_and_fixes_stale_examples():
    manual = (ROOT / RA.MANUAL_FILENAME).read_text(encoding="utf-8", errors="ignore")
    assert "set Corrupt files to Keep and report" in manual
    assert "This also makes Folder removal effectively Never" in manual
    assert "The default corruption settings are unchanged" in manual
    assert "TLO_V1.7Build<BUILD>_complete_Linux.zip" in manual
    assert "Build524" not in release_history()
    assert "tlo-main, tlo-gi, tlo-tag, tlo-deleteDupes, tlo-research, tlo-reverse, tlo-search, search-artist-db" in manual
    assert manual.count("Copies are compared with one another as well as with unsuffixed folders.") == 1
    assert "Build 532 - " in release_history()
    assert 'Release verification integrity: added the Ruff release lint gate' in release_history()
    assert 'Later releases attest Ruff only when the pinned tool actually executes.' in release_history()
    assert "its DISPLAY-unset/zero-GUI-skip wording was later corrected by Build 532" in release_history()


def test_build548_faq_describes_current_behavior_and_default_destructive_corruption_action():
    faq = (ROOT / "TLO-FAQ.txt").read_text(encoding="utf-8")
    assert "Q: Does TLO inventory absolutely every music file?\nA: No." in faq
    assert "By default, Full Inventory and standalone tlo-tag move proven-corrupt FLAC/MP3 files" in faq
    assert "Full Inventory also moves a folder when 100% of its audio files are corrupt" in faq
    assert "Keep and report" in faq
    assert "SHA-256" in faq
    assert "Q: What does Close do in Redundancy Groups?" in faq
    assert "Q: What support email is shown in Help > About?" in faq
    assert not re.search(r"\bBuild\s+\d+\b", faq)
    assert "Build 488 note" not in faq
    assert "Build 489 note" not in faq


def test_build548_test_metadata_is_current_without_per_test_version_stamps_or_current_filename_literals():
    assert not (ROOT / "tests" / "test_update_signing_v513.py").exists()
    moved = ROOT / "tests" / "unit" / "test_update_signing_v513.py"
    assert moved.is_file()
    assert "Build 513" in (ast.get_docstring(ast.parse(moved.read_text(encoding="utf-8"))) or "")

    version_offenders = []
    filename_offenders = []
    current_names = {
        RA.REQUIREMENTS_FILENAME,
        RA.MANUAL_FILENAME,
        RA.BUILD_VERIFICATION_FILENAME,
        RA.SOURCE_README_FILENAME,
    }
    candidates = list((ROOT / "tests").rglob("*.py")) + [ROOT / "test_tlo_requirements.py", ROOT / "conftest.py"]
    for path in candidates:
        text = path.read_text(encoding="utf-8", errors="replace")
        tree = ast.parse(text)
        for node in tree.body:
            if isinstance(node, (ast.Assign, ast.AnnAssign)):
                targets = node.targets if isinstance(node, ast.Assign) else [node.target]
                if any(isinstance(t, ast.Name) and t.id == "__version__" for t in targets):
                    version_offenders.append(str(path.relative_to(ROOT)))
        if path.name not in {Path(__file__).name, "_release_artifacts.py"} and any(name in text for name in current_names):
            filename_offenders.append(str(path.relative_to(ROOT)))
    assert version_offenders == []
    assert filename_offenders == []
    helper = (ROOT / "tests" / "_release_artifacts.py").read_text(encoding="utf-8")
    assert "import tlo_version as V" in helper
    assert 'f"TLO_Inventory_Requirements_Working_{V.VERSION}.docx"' in helper


def test_build548_stale_build_labels_are_removed_from_build_metadata():
    ffmpeg_source = (ROOT / "FFMPEG_BUILD_SOURCE.txt").read_text(encoding="utf-8")
    build_lock = (ROOT / "requirements-build.txt").read_text(encoding="utf-8")
    prepare = (ROOT / "prepare_ffmpeg.py").read_text(encoding="utf-8")
    assert not re.search(r"TLO Build \d+", ffmpeg_source)
    assert not re.search(r"TLO Build \d+", build_lock)
    assert "TLO-build/535" not in prepare
    assert 'f"TLO-build/{__version__.removeprefix(\'v\')}"' in prepare
