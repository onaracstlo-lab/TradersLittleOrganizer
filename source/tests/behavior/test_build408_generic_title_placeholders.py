"""Build 408 regressions for generic Title/Titled/Titles placeholders."""

from pathlib import Path

import pytest

pytestmark = pytest.mark.behavior

import tlo_tag_lib as T

__version__ = "v413"


@pytest.mark.parametrize(
    "raw",
    ["Title", "TITLE", "title", "Titled", "tItLeD", "Titles", "TiTlEs"],
)
def test_build408_existing_audio_title_tags_make_generic_title_words_unknown(raw):
    assert T._usable_title_from_audio_title_tag(raw) == ("Unknown", False)


@pytest.mark.parametrize(
    "raw",
    ["01. Title", "002 - TITLED", "3: titles"],
)
def test_build408_numbered_existing_audio_title_tags_make_generic_title_words_unknown(raw):
    assert T._usable_title_from_audio_title_tag(raw) == ("Unknown", False)


@pytest.mark.parametrize(
    "raw",
    ["The Title Track", "A Song Titled Sue", "Titles Are Hard"],
)
def test_build408_existing_audio_title_tags_do_not_reject_longer_real_titles(raw):
    assert T._usable_title_from_audio_title_tag(raw) == (raw, True)


def test_build408_numbered_setlist_rows_keep_positions_but_write_unknown(tmp_path: Path):
    setlist = tmp_path / "numbered.txt"
    setlist.write_text("1 Title\n2 tItLeD\n3 TITLES\n4 Real Song\n", encoding="utf-8")

    tracks = T.parse_setlist_tracks(str(setlist))

    assert [(row["normalized_number"], row["title"]) for row in tracks] == [
        (1, "Unknown"),
        (2, "Unknown"),
        (3, "Unknown"),
        (4, "Real Song"),
    ]


def test_build408_unnumbered_setlist_section_makes_generic_title_words_unknown(tmp_path: Path):
    setlist = tmp_path / "unnumbered-section.txt"
    setlist.write_text("Set 1\nTitle\nTITLED\ntitles\nReal Song\n", encoding="utf-8")

    tracks, source = T.parse_unnumbered_section_tracks(str(setlist), expected_count=4)

    assert source == "unnumbered-sections"
    assert [row["title"] for row in tracks] == ["Unknown", "Unknown", "Unknown", "Real Song"]


def test_build408_unstructured_setlist_block_makes_generic_title_words_unknown(tmp_path: Path):
    setlist = tmp_path / "unnumbered.txt"
    setlist.write_text("Title\nTitled\nTitles\nReal Song\n", encoding="utf-8")

    tracks, source = T.parse_unstructured_unnumbered_tracks(str(setlist), expected_count=4)

    assert source == "unnumbered-lines"
    assert [row["title"] for row in tracks] == ["Unknown", "Unknown", "Unknown", "Real Song"]


def test_build408_comma_setlist_items_make_generic_title_words_unknown(tmp_path: Path):
    setlist = tmp_path / "comma.txt"
    setlist.write_text("Alpha, Title, Titled, Titles, Omega\n", encoding="utf-8")

    tracks, source = T.parse_unnumbered_comma_tracks(str(setlist), expected_count=5)

    assert source == "comma-items"
    assert [row["title"] for row in tracks] == ["Alpha", "Unknown", "Unknown", "Unknown", "Omega"]


def test_build408_longer_setlist_titles_containing_words_remain_unchanged(tmp_path: Path):
    setlist = tmp_path / "longer.txt"
    setlist.write_text("1 The Title Track\n2 A Song Titled Sue\n3 Titles Are Hard\n", encoding="utf-8")

    tracks = T.parse_setlist_tracks(str(setlist))

    assert [row["title"] for row in tracks] == ["The Title Track", "A Song Titled Sue", "Titles Are Hard"]


def test_build408_requirements_and_manual_document_generic_title_rule():
    from docx import Document

    root = Path(__file__).resolve().parents[2]
    requirements = Document(root / "TLO_Inventory_Requirements_Working_v413.docx")
    req_text = "\n".join(p.text for p in requirements.paragraphs)
    manual_text = (root / "TLO_Inventory_User_Manual_v413.rtf").read_text(encoding="utf-8", errors="ignore")

    assert "Current document version: v413 (v1.4 Build 413)." in req_text
    assert "complete cleaned candidate song title" in req_text
    assert "exactly Title, Titled, or Titles" in req_text
    assert "keep the track row/position but write the Title value as Unknown" in req_text
    assert "Version v1.4 Build 413" in manual_text
    assert "Build 408 - Generic song-title placeholders" in manual_text
    assert "exactly Title, Titled, or Titles after normal cleanup" in manual_text
